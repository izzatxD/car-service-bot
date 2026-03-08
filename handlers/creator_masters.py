from aiogram import Router, F, types
from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton, BufferedInputFile
from aiogram.utils.keyboard import InlineKeyboardBuilder

from config import CREATOR_ID
from db.core import get_pool
from db import get_all_masters, get_master, get_daily_stats, get_monthly_stats, get_all_time_stats, get_active_orders, get_history_by_period, get_all_orders, fmt_dt
from handlers.creator_main import render_dashboard
from handlers.creator_utils import _safe_edit

router = Router()
PAGE_SIZE = 10

async def render_masters_page(page: int) -> tuple[str, InlineKeyboardMarkup]:
    masters = await get_all_masters()
    total   = len(masters)
    total_pages = max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE)
    page = max(0, min(page, total_pages - 1))

    start = page * PAGE_SIZE
    chunk = masters[start: start + PAGE_SIZE]

    active_count   = sum(1 for m in masters if m['is_active'])
    blocked_count  = total - active_count

    text = (
        f"👨‍🔧 <b>USTALAR RO'YXATI</b>  ({page + 1}/{total_pages} sahifa)\n"
        "━━━━━━━━━━━━━━━━━━\n"
        f"🟢 Faol: <b>{active_count}</b>   🔴 Bloklangan: <b>{blocked_count}</b>   Jami: <b>{total}</b>\n\n"
        "<i>Usta profiliga kirish yoki blok/faol qilish uchun tugmani bosing:</i>"
    )

    builder = InlineKeyboardBuilder()

    for m in chunk:
        status_icon = "🟢" if m['is_active'] else "🔴"
        name = m.get('full_name') or m.get('username') or f"ID:{m['telegram_id']}"
        display = f"{status_icon} {name[:22]}"
        toggle_label = "🔒 Bloklash" if m['is_active'] else "🔓 Faollashtirish"

        builder.row(
            InlineKeyboardButton(text=display,       callback_data=f"admin_master_{m['telegram_id']}"),
            InlineKeyboardButton(text=toggle_label,  callback_data=f"admin_toggle_{m['telegram_id']}_page_{page}"),
        )

    nav_buttons = []
    if page > 0:
        nav_buttons.append(InlineKeyboardButton(text="◀ Oldingi", callback_data=f"admin_masters_page_{page - 1}"))
    nav_buttons.append(InlineKeyboardButton(text=f"{page + 1}/{total_pages}", callback_data="admin_noop"))
    if page < total_pages - 1:
        nav_buttons.append(InlineKeyboardButton(text="Keyingi ▶", callback_data=f"admin_masters_page_{page + 1}"))

    builder.row(*nav_buttons)
    builder.row(
        InlineKeyboardButton(text="🔍 Qidirish",   callback_data="admin_search_master"),
        InlineKeyboardButton(text="🔙 Dashboard", callback_data="admin_dashboard"),
    )

    return text, builder.as_markup()


async def render_master_profile(master_id: int) -> tuple[str, InlineKeyboardMarkup]:
    master = await get_master(master_id)
    if not master:
        return "❌ Usta topilmadi.", InlineKeyboardMarkup(inline_keyboard=[[
            InlineKeyboardButton(text="🔙 Orqaga", callback_data="admin_masters_page_0")
        ]])

    status = "🟢 FAOL" if master['is_active'] else "🔴 BLOKLANGAN"
    name     = master.get('full_name') or "Noma'lum"
    username = f"@{master['username']}" if master.get('username') else "—"

    daily   = await get_daily_stats(master_id)
    monthly = await get_monthly_stats(master_id)
    # get_all_time_stats already imported from db at top
    all_time = await get_all_time_stats(master_id)

    text = (
        f"{status}  <b>{name}</b>  {username}\n"
        f"🆔 <code>{master_id}</code>\n"
        "━━━━━━━━━━━━━━━━━━\n"
        "📅 <b>Bugun:</b>\n"
        f"  📥 Qabul: <b>{daily['received_count']}</b>  ✅ Topshiri: <b>{daily['completed_count']}</b>\n"
        f"  💰 Tushum: <b>{daily['completed_total']:,.0f} s</b>  💵 Foyda: <b>{(daily['completed_total'] - daily['completed_cost']):,.0f} s</b>\n\n"
        "📆 <b>Shu oy:</b>\n"
        f"  📥 Qabul: <b>{monthly['received_count']}</b>  ✅ Topshiri: <b>{monthly['completed_count']}</b>\n"
        f"  📦 Faol Ishlar: <b>{monthly['active_count']}</b>\n"
        f"  💰 Tushum: <b>{monthly['completed_total']:,.0f} s</b>  💵 Foyda: <b>{(monthly['completed_total'] - monthly['completed_cost']):,.0f} s</b>\n\n"
        "📈 <b>Umumiy (Barcha vaqt):</b>\n"
        f"  👥 Jami xizmat ko'rsatilganlar: <b>{all_time['completed_count']}</b> ta mijoz\n"
        f"  💰 Umumiy ishlagan summasi: <b>{all_time['completed_total']:,.0f} s</b>\n"
        f"  💵 Umumiy foyda: <b>{(all_time['completed_total'] - all_time['completed_cost']):,.0f} s</b>\n"
    )

    toggle_label = "🔒 Bloklash" if master['is_active'] else "🔓 Faollashtirish"

    builder = InlineKeyboardBuilder()
    builder.row(InlineKeyboardButton(text="⏳ Faol Ishlari (Tugatmagan)", callback_data=f"admin_m_orders_{master_id}"))
    builder.row(InlineKeyboardButton(text="✅ Topshirgan ishlari", callback_data=f"admin_m_history_menu_{master_id}"))
    builder.row(
        InlineKeyboardButton(text=toggle_label, callback_data=f"admin_toggle_{master_id}_profile"),
        InlineKeyboardButton(text="🗑 O'chirish", callback_data=f"admin_del_ask_{master_id}"),
    )
    builder.row(InlineKeyboardButton(text="🔙 Ro'yxatga", callback_data="admin_masters_page_0"))

    return text, builder.as_markup()


async def render_active_orders(filter_type: str = "all", master_id: int = None) -> tuple[str, InlineKeyboardMarkup]:
    all_orders = await get_active_orders(master_id)

    if filter_type == "progress":
        orders = [o for o in all_orders if o['status'] in ['qabul_qilindi', 'jarayonda']]
    elif filter_type == "ready":
        orders = [o for o in all_orders if o['status'] == 'tayyor']
    else:
        orders = all_orders

    text = f"📦 <b>FAOL ISHLAR</b> ({len(orders)} ta)\n━━━━━━━━━━━━━━━━━━\n\n"

    for o in orders[:25]:
        st_icon = "🟡" if o['status'] == "jarayonda" else "🟢" if o['status'] == "tayyor" else "⚪"
        master_tag = f"<b>M:</b>{o['master_id']} | " if not master_id else ""
        st_text = o['status'].replace('_', ' ').capitalize()
        text += (
            f"{st_icon} <b>#{o['id']}</b> | {o['client_name']} | 🚗 {o['car_model']} ({o['car_number']})\n"
            f"   └ {master_tag}{st_text} | 💰 {o['price']:,.0f} so'm\n\n"
        )

    if not orders:
        text += "Ushbu filtrda faol ishlar topilmadi."

    builder = InlineKeyboardBuilder()

    f_all    = "🔘 Barchasi"  if filter_type == "all"      else "Barchasi"
    f_prog   = "🔘 Jarayonda" if filter_type == "progress" else "Jarayonda"
    f_ready  = "🔘 Tayyor"    if filter_type == "ready"    else "Tayyor"

    master_param = f"_{master_id}" if master_id else ""
    builder.row(
        InlineKeyboardButton(text=f_all,   callback_data=f"admin_orders_all{master_param}"),
        InlineKeyboardButton(text=f_prog,  callback_data=f"admin_orders_progress{master_param}"),
        InlineKeyboardButton(text=f_ready, callback_data=f"admin_orders_ready{master_param}"),
    )

    back_cb = f"admin_master_{master_id}" if master_id else "admin_dashboard"
    builder.row(InlineKeyboardButton(text="🔙 Orqaga", callback_data=back_cb))

    return text, builder.as_markup()

@router.callback_query(F.data.startswith("admin_masters_page_"))
async def cb_masters_page(cb: types.CallbackQuery):
    if cb.from_user.id != CREATOR_ID:
        return
    page = int(cb.data.split("_")[-1])
    text, markup = await render_masters_page(page)
    await _safe_edit(cb, text, markup)


@router.callback_query(F.data.startswith("admin_master_"))
async def cb_admin_master_detail(cb: types.CallbackQuery):
    if cb.from_user.id != CREATOR_ID:
        return
    master_id = int(cb.data.split("_")[-1])
    text, markup = await render_master_profile(master_id)
    await _safe_edit(cb, text, markup)


@router.callback_query(F.data.startswith("admin_toggle_"))
async def cb_admin_toggle_master(cb: types.CallbackQuery):
    if cb.from_user.id != CREATOR_ID:
        return

    parts = cb.data.split("_")
    master_id = int(parts[2])
    source    = "_".join(parts[3:])

    pool = get_pool()
    async with pool.acquire() as conn:
        await conn.execute(
            "UPDATE masters SET is_active = NOT is_active WHERE telegram_id = $1",
            master_id
        )

    if source == "profile":
        text, markup = await render_master_profile(master_id)
    elif source.startswith("page_"):
        page = int(source.split("_")[1])
        text, markup = await render_masters_page(page)
    else:
        text, markup = await render_dashboard()

    await _safe_edit(cb, text, markup)
    await cb.answer("✅ Holat yangilandi!")


@router.callback_query(F.data.startswith("admin_del_ask_"))
async def cb_admin_delete_ask(cb: types.CallbackQuery):
    if cb.from_user.id != CREATOR_ID:
        return
    master_id = int(cb.data.split("_")[-1])
    builder = InlineKeyboardBuilder()
    builder.row(
        InlineKeyboardButton(text="✅ Ha, o'chirilsin", callback_data=f"admin_del_yes_{master_id}"),
        InlineKeyboardButton(text="❌ Bekor",           callback_data=f"admin_master_{master_id}"),
    )
    await _safe_edit(
        cb,
        f"⚠️ <b>Diqqat!</b>\nID <code>{master_id}</code> li ustani o'chirmoqchimisiz?\n"
        "<i>Bu usta bazadan butunlay o'chiriladi.</i>",
        builder.as_markup()
    )


@router.callback_query(F.data.startswith("admin_del_yes_"))
async def cb_admin_delete_yes(cb: types.CallbackQuery):
    if cb.from_user.id != CREATOR_ID:
        return
    master_id = int(cb.data.split("_")[-1])
    pool = get_pool()
    async with pool.acquire() as conn:
        await conn.execute("DELETE FROM masters WHERE telegram_id = $1", master_id)
    await cb.answer("🗑 Usta o'chirildi!", show_alert=True)
    text, markup = await render_masters_page(0)
    await _safe_edit(cb, text, markup)


@router.callback_query(F.data.startswith("admin_orders_"))
async def cb_admin_orders_monitor(cb: types.CallbackQuery):
    if cb.from_user.id != CREATOR_ID:
        return

    data        = cb.data.replace("admin_orders_", "")
    parts       = data.split("_")
    filter_type = parts[0]
    master_id   = int(parts[1]) if len(parts) > 1 else None

    text, markup = await render_active_orders(filter_type, master_id)
    await _safe_edit(cb, text, markup)


@router.callback_query(F.data.startswith("admin_m_orders_"))
async def cb_admin_master_orders(cb: types.CallbackQuery):
    if cb.from_user.id != CREATOR_ID:
        return
    master_id = int(cb.data.split("_")[-1])
    text, markup = await render_active_orders("all", master_id)
    await _safe_edit(cb, text, markup)

@router.callback_query(F.data.startswith("admin_m_history_menu_"))
async def cb_admin_master_history_menu(cb: types.CallbackQuery):
    if cb.from_user.id != CREATOR_ID: return
    master_id = int(cb.data.split("_")[-1])
    
    text = "✅ <b>Ustaning topshirgan ishlari</b>\nQaysi vaqt oralig'ida ko'rmoqchisiz?"
    
    builder = InlineKeyboardBuilder()
    builder.row(
        InlineKeyboardButton(text="📅 Kunlik", callback_data=f"admin_m_hist_today_{master_id}"),
        InlineKeyboardButton(text="🗓 Haftalik", callback_data=f"admin_m_hist_weekly_{master_id}")
    )
    builder.row(
        InlineKeyboardButton(text="📆 Oylik", callback_data=f"admin_m_hist_monthly_{master_id}"),
        InlineKeyboardButton(text="📈 Umumiy", callback_data=f"admin_m_hist_all_{master_id}")
    )
    builder.row(InlineKeyboardButton(text="📥 Barchasini saqlash (Excel)", callback_data=f"admin_m_hist_excel_{master_id}"))
    builder.row(InlineKeyboardButton(text="🔙 Orqaga", callback_data=f"admin_master_{master_id}"))
    
    await _safe_edit(cb, text, builder.as_markup())

@router.callback_query(F.data.startswith("admin_m_hist_"))
async def cb_admin_m_hist(cb: types.CallbackQuery):
    if cb.from_user.id != CREATOR_ID: return
    parts = cb.data.split("_")
    period = parts[3]
    master_id = int(parts[4])
    
    if period == "excel":
        await cb.answer("⏳ Excel fayl tayyorlanmoqda...", show_alert=False)
        # It's better to reuse mechanic's excel export logic
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        import io
        orders = await get_all_orders(master_id)
        if not orders:
            return await cb.message.answer("😕 Bazada ushbu ustaning hali mijozlari yo'q.")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Mijozlar Bazasi"
        
        headers = ["T/r", "Qabul Qilingan Vaqt", "Status", "Mijoz Ismi", "Telefon", "Mashina Modeli", "Davlat Raqami", "Probeg", "Muammo", "Narx", "Xarajat", "Sof Foyda"]
        ws.append(headers)
        
        for k, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=k)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="cccccc", end_color="cccccc", fill_type="solid")
            
        for i, o in enumerate(orders, 1):
            profit = (o['price'] or 0) - (o['cost'] or 0)
            ws.append([
                i,
                o['created_at'],
                o['status'],
                o['client_name'],
                o['client_phone'],
                o['car_model'],
                o['car_number'],
                o.get('mileage', ''),
                o['problem'],
                o['price'],
                o['cost'],
                profit
            ])
            
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = min(adjusted_width, 40)
            
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        
        from datetime import datetime
        date_str = datetime.now().strftime('%Y-%m-%d')
        doc = BufferedInputFile(out.read(), filename=f"Mijozlar_Ustasi_{master_id}_{date_str}.xlsx")
        await cb.message.answer_document(doc, caption=f"📊 Ustaga tegishli barcha buyurtmalar bazasi")
        await cb.answer()
        return

    # get_history_by_period and fmt_dt already imported from db at top
    history = await get_history_by_period(master_id, period)
    
    label_map = {"today": "Kunlik", "weekly": "Haftalik", "monthly": "Oylik", "all": "Umumiy barcha"}
    label = label_map.get(period, period)
    
    if not history:
        await cb.answer(f"😕 {label} ro'yxatda hali yakunlagan ishlari yo'q.", show_alert=True)
        return
        
    text = f"✅ <b>Ustaning {label} topshirgan ishlari:</b>\n━━━━━━━━━━━━━━━━━━━━━━\n\n"
    for i, o in enumerate(history, 1):
        # Limit to 50 exactly in telegram text, rest user can see in excel
        if i > 50:
            break
        text += (
            f"<b>{i}) {o['client_name']}</b>\n"
            f"   🚗 {o['car_model']} ({o.get('car_number') or '-'}) | 🔧 {o['problem']}\n"
            f"   💰 {o['price']:,.0f} so'm | 📅 {fmt_dt(o['created_at'])}\n\n"
        )
    
    if len(history) > 50:
        text += f"\n... jami {len(history)} ta ish. Hammasini ko'rish uchun Excel formatda yuklab oling.\n"
    elif len(text) > 4000:
        text = text[:4000] + "\n... (davomi bor)"
        
    builder = InlineKeyboardBuilder()
    builder.row(InlineKeyboardButton(text="🔙 Orqaga", callback_data=f"admin_m_history_menu_{master_id}"))
    
    await _safe_edit(cb, text, builder.as_markup())

