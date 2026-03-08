import logging
import os
from datetime import datetime
from aiogram import Router, F, types
from aiogram.types import InlineKeyboardMarkup, InlineKeyboardButton, FSInputFile

from config import CREATOR_ID
from db import get_master, get_branches, get_daily_stats, get_monthly_stats, get_weekly_history, get_all_orders, fmt_dt

logger = logging.getLogger(__name__)
router = Router()

def build_branch_keyboard(branches: list[str], prefix: str):
    buttons = []
    for i in range(0, len(branches), 2):
        row = [InlineKeyboardButton(text=f"🏢 {branches[i]}", callback_data=f"{prefix}:{branches[i]}")]
        if i + 1 < len(branches):
            row.append(InlineKeyboardButton(text=f"🏢 {branches[i+1]}", callback_data=f"{prefix}:{branches[i+1]}"))
        buttons.append(row)
    buttons.append([InlineKeyboardButton(text="📊 Umumiy (Barchasi)", callback_data=f"{prefix}:all")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)

# ═══════════════════════════════════════════════
#  📊 STATISTIKA
# ═══════════════════════════════════════════════
@router.message(F.text == "📊 Statistika")
async def stats_handler(message: types.Message):
    master_id = message.from_user.id # Har doim o'ziniki
    master = await get_master(master_id)
    if master and master.get('role') == 'boss':
        branches = await get_branches(master_id)
        if branches:
            await message.answer(
                "📊 <b>Statistika</b> tahlili uchun qaysi filialni tanlaysiz?",
                reply_markup=build_branch_keyboard(branches, "stats_branch")
            )
            return
            
    await _show_statistics(message, master_id, None)

@router.callback_query(F.data.startswith("stats_branch:"))
async def stats_branch_callback(callback: types.CallbackQuery):
    bn = callback.data.split(":", 1)[1]
    branch = None if bn == "all" else bn
    master_id = callback.from_user.id # Har doim o'ziniki
    await callback.message.delete()
    await _show_statistics(callback.message, master_id, branch, is_edit=True)
    await callback.answer()

async def _show_statistics(message: types.Message, master_id: int | None, branch_name: str | None, is_edit: bool = False):
    daily   = await get_daily_stats(master_id, branch_name)
    monthly = await get_monthly_stats(master_id, branch_name)
    
    d_profit  = daily['completed_total']  - daily['completed_cost']
    m_profit  = monthly['completed_total'] - monthly['completed_cost']
    
    bn_text = f" ({branch_name})" if branch_name else ""

    text = (
        f"📊 <b>STATISTIKA</b>{bn_text}\n"
        "━━━━━━━━━━━━━━━━━━\n"
        "📅 <b>Bugun:</b>\n"
        f"  📥 Qabul: <b>{daily['received_count']}</b>  ✅ Topshirildi: <b>{daily['completed_count']}</b>\n"
        f"  💰 Tushum: <b>{daily['completed_total']:,.0f} so'm</b>\n"
        f"  📉 Harajat: {daily['completed_cost']:,.0f} so'm\n"
        f"  💵 Sof Foyda: <b>{d_profit:,.0f} so'm</b>\n"
        f"  💳 Naqd: {daily['naqd']:,.0f}  |  Karta: {daily['karta']:,.0f}  |  Nasiya: {daily['nasiya']:,.0f}\n\n"
        "📆 <b>Shu oy:</b>\n"
        f"  📥 Qabul: <b>{monthly['received_count']}</b>  ✅ Topshirildi: <b>{monthly['completed_count']}</b>\n"
        f"  📦 Faol: <b>{monthly.get('active_count', 0)}</b>\n"
        f"  💰 Tushum: <b>{monthly['completed_total']:,.0f} so'm</b>\n"
        f"  📉 Harajat: {monthly['completed_cost']:,.0f} so'm\n"
        f"  💵 Sof Foyda: <b>{m_profit:,.0f} so'm</b>\n"
        f"  💳 Naqd: {monthly['naqd']:,.0f}  |  Karta: {monthly['karta']:,.0f}  |  Nasiya: {monthly['nasiya']:,.0f}\n"
    )
    
    bn_data = branch_name if branch_name else "all"
    markup = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🗓 Oxirgi 7 kunlik tarix", callback_data=f"stats_weekly_{bn_data}")],
        [InlineKeyboardButton(text="📥 Barcha Mijozlar (Excel)", callback_data=f"stats_excel_{bn_data}")]
    ])
    
    if is_edit:
        await message.edit_text(text, reply_markup=markup)
    else:
        await message.answer(text, reply_markup=markup)

@router.callback_query(F.data.startswith("stats_weekly_"))
async def process_weekly_history(callback: types.CallbackQuery):
    bn = callback.data.replace("stats_weekly_", "")
    branch_name = None if bn == "all" else bn
    master_id = callback.from_user.id # Har doim o'ziniki
    history = await get_weekly_history(master_id, branch_name)
    
    if not history:
        await callback.answer("😕 Oxirgi 7 kunda hech qanday tarix topilmadi.", show_alert=True)
        return
        
    bn_text = f" ({branch_name})" if branch_name else ""
    text = f"🗓 <b>Oxirgi 7 kunlik tarix{bn_text}:</b>\n━━━━━━━━━━━━━━━━━━━━━━\n\n"
    for i, o in enumerate(history, 1):
        st = o['status'].upper()
        if st == "TOP_SHIRILDI": st = "TOPSHIRILDI 📦"
        text += (
            f"<b>{i}. {o['client_name']} ({st})</b>\n"
            f"   🚗 {o['car_model']} ({o.get('car_number') or '-'}) | 🔧 {o['problem']}\n"
            f"   💰 {o['price']:,.0f} so'm | 📅 {fmt_dt(o['created_at'])}\n\n"
        )
    
    if len(text) > 4000:
        text = text[:4000] + "\n... (davomi bor)"
        
    await callback.message.answer(text)
    await callback.answer()

@router.callback_query(F.data.startswith("stats_excel_"))
async def process_excel_export(callback: types.CallbackQuery):
    await callback.answer("⏳ Excel fayl tayyorlanmoqda...", show_alert=False)
    bn = callback.data.replace("stats_excel_", "")
    branch_name = None if bn == "all" else bn
    master_id = callback.from_user.id # Har doim o'ziniki
    
    orders = await get_all_orders(master_id, branch_name)
    if not orders:
        return await callback.message.answer("😕 Bazada hali mijozlar yo'q.")
        
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Mijozlar Bazasi{' - ' + branch_name if branch_name else ''}"
    
    headers = ["T/r", "Qabul Qilingan Vaqt", "Status", "Mijoz Ismi", "Telefon", "Mashina Modeli", "Davlat Raqami", "Probeg", "Muammo", "Narx", "Xarajat", "Sof Foyda"]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="10b981", end_color="10b981", fill_type="solid")
    header_font = Font(bold=True, color="ffffff")
    for col_num, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    total_price = 0
    total_cost = 0
    total_profit = 0
        
    for i, o in enumerate(orders, 1):
        dt_val = o.get('created_at', '')[:16]
        price = o.get('price') or 0
        cost = o.get('cost') or 0
        profit = price - cost
        
        total_price += price
        total_cost += cost
        total_profit += profit
        
        row = [
            i,
            dt_val,
            str(o.get('status', '')).upper(),
            o.get('client_name', ''),
            str(o.get('client_phone', '')),
            o.get('car_model', ''),
            o.get('car_number', ''),
            o.get('mileage', ''),
            o.get('problem', ''),
            price,
            cost,
            profit
        ]
        ws.append(row)
        
    # Yig'indini (Jami) jadval oxiriga yozish
    ws.append([])
    ws.append(["", "", "", "", "", "", "", "", "JAMI YIG'INDI:", total_price, total_cost, total_profit])
    
    total_row_idx = ws.max_row
    # Jami qatorini qalin qilib belgilash
    for col_idx in range(9, 13):
        cell = ws.cell(row=total_row_idx, column=col_idx)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="e6f2ff", end_color="e6f2ff", fill_type="solid")
        
    # Auto-adjust column widths roughly
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = min(adjusted_width, 40)
        
    file_name = f"Mijozlar_Bazasi_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    wb.save(file_name)
    
    await callback.message.answer_document(
        document=FSInputFile(file_name),
        caption="📊 Barcha mijozlaringiz ro'yxati (Excel formatida)"
    )
    
    # Qoldiqni o'chirish
    try:
        os.remove(file_name)
    except:
        pass
    
    await callback.answer()
